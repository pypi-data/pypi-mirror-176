import os
from pathlib import Path
from typing import Generator

from docx import Document as DocumentLoader
from docx.document import Document
from docx.text.paragraph import Paragraph
from docx.table import _Row, _Cell

from openpyxl import load_workbook

from ..errors import NotImplemented, InconsistentData
from ..functions import tokens_bind

from .__base import BaseScript


class Generate(BaseScript):
    def __init__(
        self,
        input: Path,
        input_format: str,
        template: Path,
        output: Path,
        document_name_column: str,
    ) -> None:
        super().__init__()
        self.input = input
        self.input_format = input_format
        self.template = template
        self.output = output
        self.document_name_column = document_name_column

    def read_excel(self):
        input_absolute = self.input.absolute()
        input_str = str(input_absolute)
        wb = load_workbook(input_str, data_only=True)
        fws = wb.worksheets[0]
        headers = None
        headers_len = 0
        data = []

        for row in fws.values:
            if not headers:
                headers = row
                headers_len = len(row)
                continue  # to next row

            if len(row) > headers_len:
                raise InconsistentData("A row has more cells than headers available")

            row_dict = dict()

            for i, cell_value in enumerate(row):
                key = headers[i]
                key_lowercase = key.lower()
                row_dict[key_lowercase] = cell_value

            data.append(row_dict)

        return data

    def get_cells(self, doc: Document) -> Generator[_Cell, None, None]:
        for table in doc.tables:
            row: _Row
            for row in table.rows:
                cell: _Cell
                for cell in row.cells:
                    yield cell

    def get_paragraphs(self, doc: Document) -> Generator[Paragraph, None, None]:
        for paragraph in doc.paragraphs:
            yield paragraph

    def generate_document(self, data: dict):
        template_absolute = self.template.absolute()
        template_str = str(template_absolute)
        doc: Document = DocumentLoader(template_str)

        cells = self.get_cells(doc)
        for cell in cells:
            cell.text = tokens_bind(cell.text, repl_map=data)

        paragraphs = self.get_paragraphs(doc)
        for paragraph in paragraphs:
            tokenized_text = tokens_bind(paragraph.text, repl_map=data)
            if tokenized_text != paragraph.text:
                paragraph.text = tokenized_text

        if not self.output.exists():
            os.makedirs(self.output)

        filename: str = data.get(self.document_name_column.lower())

        if not filename:
            raise InconsistentData(
                f"Filename is not defined within '{self.document_name_column}' column."
            )

        target_file_path = self.output / (filename + ".docx")
        target_file_path_absolute = target_file_path.absolute()
        target_file_path_str = str(target_file_path_absolute)

        doc.save(target_file_path_str)

    def run(self):
        # import json
        # print(json.dumps(self.read_excel(), indent=4))
        if self.input_format == "xlsx":
            input_data = self.read_excel()
        else:
            raise NotImplemented(
                f"input_format '{self.input_format}' is not implemented yet"
            )

        for data in input_data:
            self.generate_document(data)
